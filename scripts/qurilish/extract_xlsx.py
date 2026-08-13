# -*- coding: utf-8 -*-
"""
2026_БАРЧА_ДАСТУР xlsx -> xom CSV ekstraktori (qurilish domeni ETL, 1-qadam).

MUHIM: bu skript FAQAT xom ma'lumotni chiqaradi. Hech qanday biznes qoidasi
(SOATO, tashkilot normalizatsiyasi, muddat parseri, soha tasnifi, bosqich
bayroqlari) BU YERDA YO'Q — ularning barchasi PHP `Support` sinflarida va
test bilan qoplangan. Skript almashsa ham qoidalar o'zgarmaydi.

Qator klassifikatsiyasi:
  A ustunida raqam bor      -> obyekt qatori
  F to'lgan, A bo'sh        -> guruh sarlavhasi
                               (ish turi ro'yxatidagi qiymat -> work_type_label,
                                aks holda -> group_label; ПҚ-393 da soha,
                                boshqalarda tuman)

Ishga tushirish (Windows):
  set PYTHONIOENCODING=utf-8
  python scripts/qurilish/extract_xlsx.py [--src <xlsx>] [--out <dir>]
"""

from __future__ import annotations

import argparse
import csv
import io
import os
import sys

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

DEFAULT_SRC = "D:/kadr/2026_БАРЧА_ДАСТУР_12.08.26.xlsx"
DEFAULT_OUT = "storage/app/import/qurilish"

# Varaq -> dastur kodi (qurilish.programs.code bilan bir xil).
SHEETS = {
    "ПҚ-393": "pq393",
    "ДРАЙВЕР": "drayver",
    "ОПЕН": "open",
    "ОҒИР ТУМАН": "ogir_tuman",
    "ОҒИР МФЙ": "ogir_mfy",
    "ЯНГИ.ЎЗБ.ТУМАН": "yangi_uzb_tuman",
    "ЯНГИ.ЎЗБ.МФЙ": "yangi_uzb_mfy",
    "33 ТА ДХШ": "dxsh",
}

# Guruh sarlavhasi ish turini bildirsa — group_label emas, work_type_label bo'ladi.
WORK_TYPE_LABELS = {
    "Янги қуриш",
    "Реконструкция",
    "Мукаммал таъмирлаш",
    "Капитал таъмирлаш",
    "Жорий таъмирлаш",
}

OBJECT_HEADER = [
    "sheet", "row", "program_code", "seq", "external_id", "c_value",
    "designer_raw", "customer_raw", "name", "deadline_raw",
    "limit_amount", "carryover", "new_start",
    "f_L", "f_M", "f_N", "f_O",
    "f_P", "f_Q", "f_R",
    "f_S", "f_T", "f_U", "f_V",
    "f_W", "f_X", "f_Y", "f_Z", "f_AA", "f_AB",
    "f_AC", "f_AD", "f_AE", "f_AF", "tender_amount", "f_AH",
    "contract_count", "contract_amount", "disbursed", "disbursed_pct",
    "financed", "financed_pct",
    "handover_plan", "handover_actual", "handover_left", "overdue_flag",
    "contractor_raw", "note", "group_label", "work_type_label",
]

MONTHLY_HEADER = ["external_key", "month", "planned_amount"]


def col(letter: str) -> int:
    """Ustun harfi -> 0-asosli indeks."""
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


# Obyekt ustunlari — xlsx harfi tartibida (A..AT).
OBJ_COLS = [
    "B", "C", "D", "E", "F", "G", "I", "J", "K",
    "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V",
    "W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH",
    "AI", "AJ", "AK", "AL", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT",
]
MONTH_COLS = ["AV", "AW", "AX", "AY", "AZ", "BA", "BB", "BC", "BD", "BE", "BF", "BG"]


def cell(row: tuple, letter: str):
    i = col(letter)
    return row[i] if i < len(row) else None


def clean(v) -> str:
    """Matnni tozalaydi: None -> '', ichki qator uzilishi va ortiqcha bo'shliq -> bitta bo'shliq."""
    if v is None:
        return ""
    s = str(v).replace("\n", " ").replace("\r", " ").strip()
    return " ".join(s.split())


def num(v) -> str:
    """Raqamni CSV uchun: son bo'lmasa bo'sh qator."""
    return "" if not isinstance(v, (int, float)) or isinstance(v, bool) else repr(float(v))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=DEFAULT_SRC)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()

    if not os.path.isfile(args.src):
        print(f"XATO: manba topilmadi: {args.src}", file=sys.stderr)
        return 1

    os.makedirs(args.out, exist_ok=True)
    wb = openpyxl.load_workbook(args.src, read_only=True, data_only=True)

    objects_path = os.path.join(args.out, "objects.csv")
    monthly_path = os.path.join(args.out, "monthly.csv")

    n_obj = 0
    n_month = 0

    with open(objects_path, "w", encoding="utf-8", newline="") as fo, \
            open(monthly_path, "w", encoding="utf-8", newline="") as fm:
        wo = csv.writer(fo)
        wm = csv.writer(fm)
        wo.writerow(OBJECT_HEADER)
        wm.writerow(MONTHLY_HEADER)

        for sheet_name, program_code in SHEETS.items():
            if sheet_name not in wb.sheetnames:
                print(f"OGOHLANTIRISH: varaq yo'q: {sheet_name}", file=sys.stderr)
                continue

            ws = wb[sheet_name]
            group_label = ""
            work_type_label = ""
            sheet_objects = 0

            # 5-qatordan boshlanadi: 1-3 sarlavha, 4 ost-sarlavha.
            for r, row in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
                a = row[0] if row else None
                f = cell(row, "F")

                if isinstance(a, (int, float)) and not isinstance(a, bool):
                    # --- Obyekt qatori ---
                    external_id = clean(cell(row, "B"))
                    # Ba'zi ID lar oxirida nuqta bilan ('...001.') — faqat raqam qoldiramiz.
                    external_id = "".join(ch for ch in external_id if ch.isdigit())
                    external_key = external_id or f"{program_code}:{r}"

                    out = [sheet_name, r, program_code, int(a), external_id]
                    out.append(clean(cell(row, "C")))         # c_value
                    out.append(clean(cell(row, "D")))         # designer_raw
                    out.append(clean(cell(row, "E")))         # customer_raw
                    out.append(clean(cell(row, "F")))         # name
                    out.append(clean(cell(row, "G")))         # deadline_raw
                    for letter in ["I", "J", "K",
                                   "L", "M", "N", "O", "P", "Q", "R",
                                   "S", "T", "U", "V",
                                   "W", "X", "Y", "Z", "AA", "AB",
                                   "AC", "AD", "AE", "AF", "AG", "AH",
                                   "AI", "AJ", "AK", "AL", "AM", "AN",
                                   "AO", "AP", "AQ", "AR"]:
                        out.append(num(cell(row, letter)))
                    out.append(clean(cell(row, "AS")))        # contractor_raw
                    out.append(clean(cell(row, "AT")))        # note
                    out.append(group_label)
                    out.append(work_type_label)
                    wo.writerow(out)
                    n_obj += 1
                    sheet_objects += 1

                    # Oylik grafik — faqat ПҚ-393 va ДРАЙВЕР varaqlarida mavjud.
                    for m, letter in enumerate(MONTH_COLS, 1):
                        v = cell(row, letter)
                        if isinstance(v, (int, float)) and not isinstance(v, bool) and v:
                            wm.writerow([external_key, m, repr(float(v))])
                            n_month += 1

                elif f is not None and clean(f):
                    # --- Guruh sarlavhasi ---
                    label = clean(f)
                    if label in WORK_TYPE_LABELS:
                        work_type_label = label
                    else:
                        group_label = label
                        work_type_label = ""

            print(f"{sheet_name:16s} -> {sheet_objects} obyekt")

    wb.close()
    print(f"\nJAMI: {n_obj} obyekt, {n_month} oylik qator")
    print(f"  {objects_path}")
    print(f"  {monthly_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
